import streamlit as st
import pandas as pd
import smtplib
import ssl
import time
import re
import sqlite3
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders
from io import BytesIO
from docx import Document  # for reading Word files

# ---------------------------
# License / Watermark
# ---------------------------
WATERMARK = "⚡ SmartMailer | © 2025 Arpan Ari | Licensed under CC BY-NC 4.0"

# ---------------------------
# DB Setup for Users
# ---------------------------
def init_db():
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS users
                 (username TEXT PRIMARY KEY, password TEXT, role TEXT)''')
    # Insert default admin if not exists
    c.execute("SELECT * FROM users WHERE username=?", ("admin",))
    if not c.fetchone():
        c.execute("INSERT INTO users VALUES (?, ?, ?)", ("admin", "admin123", "admin"))
    conn.commit()
    conn.close()

def add_user(username, password, role="user"):
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    try:
        c.execute("INSERT INTO users VALUES (?, ?, ?)", (username, password, role))
        conn.commit()
    except sqlite3.IntegrityError:
        pass
    conn.close()

def get_users():
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute("SELECT username, role FROM users")
    users = c.fetchall()
    conn.close()
    return users

def remove_user(username):
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute("DELETE FROM users WHERE username=?", (username,))
    conn.commit()
    conn.close()

def authenticate(username, password):
    conn = sqlite3.connect("users.db")
    c = conn.cursor()
    c.execute("SELECT role FROM users WHERE username=? AND password=?", (username, password))
    result = c.fetchone()
    conn.close()
    return result[0] if result else None

# ---------------------------
# INIT DB
# ---------------------------
init_db()

# ---------------------------
# Login System
# ---------------------------
st.set_page_config(page_title="📧 Internship Bulk Mailer", page_icon="📧", layout="centered")

if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.role = None
    st.session_state.username = None

if not st.session_state.logged_in:
    st.title("🔐 SmartMailer Login")
    username = st.text_input("Username")
    password = st.text_input("Password", type="password")
    if st.button("Login"):
        role = authenticate(username, password)
        if role:
            st.session_state.logged_in = True
            st.session_state.role = role
            st.session_state.username = username
            st.success(f"✅ Logged in as {role.upper()}")
            st.rerun()
        else:
            st.error("❌ Invalid credentials")
    st.stop()

# ---------------------------
# Admin Panel
# ---------------------------
if st.session_state.role == "admin":
    st.sidebar.title("👑 Admin Panel")
    st.sidebar.subheader("Manage Users")

    users = get_users()
    st.sidebar.write("Registered Users:")
    for u in users:
        st.sidebar.write(f"- {u[0]} ({u[1]})")

    new_user = st.sidebar.text_input("New Username")
    new_pass = st.sidebar.text_input("New Password", type="password")
    if st.sidebar.button("Add User"):
        add_user(new_user, new_pass, "user")
        st.sidebar.success("✅ User added")
        st.rerun()

    remove_username = st.sidebar.selectbox("Remove User", [u[0] for u in users if u[0] != "admin"])
    if st.sidebar.button("Remove Selected User"):
        remove_user(remove_username)
        st.sidebar.warning(f"🗑️ Removed user: {remove_username}")
        st.rerun()

# ---------------------------
# Original Bulk Mailer Code (Unchanged)
# ---------------------------
st.title("📧 Internship Application Bulk Mailer")
st.write("Send **individual internship emails** to many recipients at once — with live status, detailed results, and a downloadable log.")

# Sidebar: SMTP & Options
with st.sidebar:
    st.header("⚙️ SMTP Settings")
    provider = st.selectbox(
        "Email Provider",
        ["Gmail (smtp.gmail.com:465)", "Outlook/Office365 (smtp.office365.com:587)", "Custom"]
    )

    if provider.startswith("Gmail"):
        smtp_host, smtp_port, use_ssl = "smtp.gmail.com", 465, True
        st.caption("📌 Gmail requires an **App Password** (2-Step Verification ➜ App Passwords).")
    elif provider.startswith("Outlook"):
        smtp_host, smtp_port, use_ssl = "smtp.office365.com", 587, False
        st.caption("📌 Outlook requires an **App Password** (if personal) or SMTP AUTH enabled by admin (if Office365 org).")
    else:
        smtp_host = st.text_input("SMTP Host", value="smtp.example.com")
        smtp_port = st.number_input("SMTP Port", value=465, step=1)
        use_ssl = st.checkbox("Use SSL", value=True)

    st.divider()
    st.header("⏱️ Sending Options")
    throttle = st.slider("Delay between emails (seconds)", 0.0, 3.0, 0.2, 0.1)
    test_mode = st.checkbox("Test mode (send only to the first 1 recipient)", value=False)

# Upload & Extract Emails
st.subheader("1) Upload Recipient List")
uploaded = st.file_uploader("Upload Excel, CSV, or Word file with emails", type=["xlsx", "xls", "csv", "docx"])
df = None

if uploaded:
    try:
        if uploaded.name.lower().endswith(".csv"):
            df = pd.read_csv(uploaded)
        elif uploaded.name.lower().endswith((".xlsx", ".xls")):
            df = pd.read_excel(uploaded)
        elif uploaded.name.lower().endswith(".docx"):
            doc = Document(uploaded)
            text = "\n".join([p.text for p in doc.paragraphs])
            emails = re.findall(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-z]{2,}", text)
            df = pd.DataFrame({"Email": list(set(emails))})
        st.success("✅ File loaded successfully.")
        st.dataframe(df.head(), use_container_width=True)
    except Exception as e:
        st.error(f"❌ Could not read file: {e}")

# Map Columns
email_col = name_col = None
if df is not None:
    email_col = st.selectbox("Select the Email column", df.columns, index=0)
    possible_name_cols = [c for c in df.columns if "name" in c.lower()] or list(df.columns)
    name_col = st.selectbox("Optional: Name column for personalization", ["(None)"] + possible_name_cols)

# Sender & Message
sender_email = st.text_input("Your Email Address")
sender_password = st.text_input("Your Email Password / App Password", type="password")
subject = st.text_input("Email Subject", value="Application for 6-Month Internship (Final Semester)")
content_type = st.radio("Email Type", ["Simple (Plain Text)", "Styled (HTML)"], index=1)

default_html = """\
<div style="font-family:Segoe UI, Arial, sans-serif; line-height:1.55; font-size:16px;">
  <p>Dear {name_or_sirmadam},</p>
  <p>
    I am writing to express my interest in a <b>6-month internship</b> (8th semester, final year)
    under your guidance. I’m keen to contribute to ongoing research/industry projects in
    <b>Computer Vision / Deep Learning / AI</b> and learn rigorously.
  </p>
  <p>
    <b>Brief profile:</b><br/>
    • Final-year B.Tech student (8th semester)<br/>
    • Experience with YOLOv8, CRNN-based OCR, Streamlit, Flask<br/>
    • Built ANPR, dashboards, and ML pipelines
  </p>
  <p>
    I would be grateful for the opportunity to discuss how I could assist your work.
    Please let me know if I can share a resume/portfolio or schedule a short call.
  </p>
  <p>Thank you for your time and consideration.</p>
  <p>Regards,<br/>
  {your_name}<br/>
  {your_phone}<br/>
  {your_university}</p>
</div>
"""

default_text = """\
Dear {name_or_sirmadam},

I am writing to express my interest in a 6-month internship (8th semester, final year)
under your guidance. I’m keen to contribute to ongoing research/industry projects in
Computer Vision / Deep Learning / AI and learn rigorously.

Brief profile:
• Final-year B.Tech student (8th semester)
• Experience with YOLOv8, CRNN-based OCR, Streamlit, Flask
• Built ANPR, dashboards, and ML pipelines

I would be grateful for the opportunity to discuss how I could assist your work.
Please let me know if I can share a resume/portfolio or schedule a short call.

Thank you for your time and consideration.

Regards,
{your_name}
{your_phone}
{your_university}
"""

body = st.text_area(
    "Message Body (use placeholders: {name_or_sirmadam}, {your_name}, {your_phone}, {your_university})",
    value=default_html if content_type == "Styled (HTML)" else default_text,
    height=260
)

# Attachment Upload
st.subheader("2) Upload Attachment (Optional)")
attachment_file = st.file_uploader("Upload your CV / Resume (PDF, DOCX, etc.)", type=["pdf", "docx", "txt", "jpg", "png"], accept_multiple_files=False)

# Helper Functions
def build_message(sender, recipient, subject, body_str, html=False, attachment=None):
    msg = MIMEMultipart()
    msg["From"] = sender
    msg["To"] = recipient
    msg["Subject"] = subject

    # Attach body
    if html:
        msg.attach(MIMEText(body_str, "html"))
    else:
        msg.attach(MIMEText(body_str, "plain"))

    # Attach file if present
    if attachment:
        part = MIMEBase("application", "octet-stream")
        part.set_payload(attachment.read())
        encoders.encode_base64(part)
        part.add_header(
            "Content-Disposition",
            f"attachment; filename={attachment.name}",
        )
        msg.attach(part)

    return msg

def send_one(server, sender, recipient, subject, body_str, html=False, attachment=None):
    msg = build_message(sender, recipient, subject, body_str, html, attachment)
    server.sendmail(sender, [recipient], msg.as_string())

def connect_smtp(host, port, use_ssl, sender, password):
    if use_ssl:  # Gmail
        context = ssl.create_default_context()
        server = smtplib.SMTP_SSL(host, port, context=context, timeout=30)
        server.login(sender, password)
        return server
    else:  # Outlook
        server = smtplib.SMTP(host, port, timeout=30)
        server.starttls(context=ssl.create_default_context())
        server.login(sender, password)
        return server

def make_log_download(df_log):
    out = BytesIO()
    df_log.to_csv(out, index=False, encoding="utf-8")
    out.seek(0)
    return out

# Preview
if df is not None:
    preview_name = "Sir/Madam"
    if name_col and name_col != "(None)" and name_col in df.columns and not df[name_col].dropna().empty:
        preview_name = str(df[name_col].dropna().iloc[0])

    preview_body = body.format(
        name_or_sirmadam=preview_name,
        your_name="Your Name",
        your_phone="+91-XXXXXXXXXX",
        your_university="Your College/University"
    )

    st.subheader("Preview Email")
    if content_type == "Styled (HTML)":
        st.markdown(preview_body, unsafe_allow_html=True)
    else:
        st.code(preview_body)

    if attachment_file:
        st.info(f"📎 Attachment selected: {attachment_file.name}")

# Send Emails
if st.button("🚀 Send Emails"):
    if df is None:
        st.error("Please upload a recipients file.")
    elif not sender_email or not sender_password or not subject or not body:
        st.error("Please complete all required fields.")
    else:
        recipients = (
            df[email_col].astype(str).str.strip().replace({"nan": None, "": None}).dropna().tolist()
        )
        if not recipients:
            st.error("No valid email addresses found.")
        else:
            if test_mode:
                recipients = recipients[:1]

            names = {}
            if name_col and name_col != "(None)" and name_col in df.columns:
                for _, row in df.iterrows():
                    em = str(row[email_col]).strip()
                    nm = str(row[name_col]).strip() if pd.notna(row.get(name_col, "")) else ""
                    if em:
                        names[em] = nm if nm else "Sir/Madam"

            results = []
            progress = st.progress(0)
            try:
                server = connect_smtp(smtp_host, smtp_port, use_ssl, sender_email, sender_password)
            except Exception as e:
                st.error(f"❌ SMTP connection/login failed: {e}")
                st.stop()

            success_count = 0
            fail_count = 0
            for idx, rcpt in enumerate(recipients, start=1):
                nm = names.get(rcpt, "Sir/Madam")
                filled_body = body.format(
                    name_or_sirmadam=nm,
                    your_name="Your Name",
                    your_phone="+91-XXXXXXXXXX",
                    your_university="Your College/University"
                )
                try:
                    send_one(server, sender_email, rcpt, subject, filled_body, html=(content_type == "Styled (HTML)"), attachment=attachment_file)
                    success_count += 1
                    results.append({"email": rcpt, "name": nm, "status": "✅ Sent", "error": ""})
                except Exception as ex:
                    fail_count += 1
                    results.append({"email": rcpt, "name": nm, "status": "❌ Failed", "error": str(ex)})
                progress.progress(int(idx * 100 / len(recipients)))
                if throttle > 0:
                    time.sleep(throttle)

            try:
                server.quit()
            except Exception:
                pass

            df_log = pd.DataFrame(results)
            st.success(f"✅ Done! Sent: {success_count} • ❌ Failed: {fail_count} • Total: {len(recipients)}")

            buf = make_log_download(df_log)
            st.download_button(
                "📥 Download Results (CSV)",
                data=buf,
                file_name="bulk_mail_results.csv",
                mime="text/csv"
            )

# ---------------------------
# Footer Watermark
# ---------------------------
st.markdown(f"<hr><center>{WATERMARK}</center>", unsafe_allow_html=True)
